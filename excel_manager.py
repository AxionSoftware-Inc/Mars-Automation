import os
import re
import logging
from datetime import datetime
from zoneinfo import ZoneInfo
from typing import Optional, Dict, Any, List
import openpyxl

logger = logging.getLogger(__name__)

DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
LATEST_EXCEL_PATH = os.path.join(DATA_DIR, "latest.xlsx")

def _normalize_hex(color_val) -> str:
    """openpyxl RGB/ARGB rangini 6 xonali standart hex formatga keltiradi."""
    if not color_val:
        return ""
    s = str(color_val).upper().strip()
    if len(s) == 8:
        # 8 xonali ARGB: alfa kanalni kesib tashlash (FF000000 -> 000000 [qora], FFFF0000 -> FF0000 [qizil])
        return s[2:]
    return s

def _is_red_color(color_val) -> bool:
    """Qizil rangni aniqlash (FF000000 kabi qora ranglar bilan chalkashmasligi shart)."""
    c = _normalize_hex(color_val)
    return c in ("FF0000", "C00000", "E74C3C", "ED1C24")

def _is_yellow_color(color_val) -> bool:
    """Sariq rangni aniqlash (P3 Humo apparat)."""
    c = _normalize_hex(color_val)
    return c in ("FFFF00", "FFFFE0", "FFF200")


class ExcelManager:
    def __init__(self, excel_path: str = LATEST_EXCEL_PATH):
        self.excel_path = excel_path
        # Xotiradagi indekslar: tezkor qidiruv uchun lug'at (dict)
        # key: stir yoki jshshir yoki terminal -> list of records
        self.stir_index: Dict[str, List[Dict[str, Any]]] = {}
        self.jshshir_index: Dict[str, List[Dict[str, Any]]] = {}
        self.terminal_index: Dict[str, List[Dict[str, Any]]] = {}
        self.phone_index: Dict[str, List[Dict[str, Any]]] = {}
        # Hisobdan chiqarilgan (qizil qator) yozuvlar indeksi:
        self.decommissioned_stir_index: Dict[str, List[Dict[str, Any]]] = {}
        self.decommissioned_jshshir_index: Dict[str, List[Dict[str, Any]]] = {}
        self.decommissioned_terminal_index: Dict[str, List[Dict[str, Any]]] = {}
        self.records_count = 0
        self.decommissioned_count = 0
        self.load_data()

    def clean_digits(self, val: Any) -> str:
        """Faqat raqamlarni ajratib olish (float / string e'tiborga olingan)"""
        if val is None:
            return ""
        s = str(val).split(".")[0]  # Masalan 311790999.0 bo'lib qolmasligi uchun
        return re.sub(r"\D", "", s)

    def load_data(self) -> bool:
        """Excel faylni xotiraga tezkor yuklash (read_only + data_only)"""
        if not os.path.exists(self.excel_path):
            logger.warning(f"Excel fayli topilmadi: {self.excel_path}")
            return False

        try:
            wb = openpyxl.load_workbook(self.excel_path, read_only=True, data_only=True)
            sheet_name = wb.sheetnames[0]
            sheet = wb[sheet_name]

            # Indekslarni tozalash
            new_stir = {}
            new_jshshir = {}
            new_terminal = {}
            new_phone = {}
            new_decom_stir = {}
            new_decom_jshshir = {}
            new_decom_terminal = {}
            count = 0
            p3_count = 0
            decom_count = 0

            MONTHS = ["Yanvar", "Fevral", "Mart", "Aprel", "May", "Iyun", "Iyul", "Avgust", "Sentyabr", "Oktyabr", "Noyabr", "Dekabr"]

            # 3-qatordan boshlab o'qish (1 va 2-qator sarlavhalar)
            # - P3 Humo apparatlar Excelda sariq qator (FFFFFF00), Smart One apparatlar esa oq qator bilan belgilangan.
            # - Qizil qatorlar (FFFF0000 yoki C00000) hisobdan chiqarilgan / o'chirib tashlangan deb hisoblanadi va hisoblanmaydi!
            for row_cells in sheet.iter_rows(min_row=3, values_only=False):
                row = [c.value if c is not None else None for c in row_cells]
                is_red = False
                is_p3 = False
                for c in row_cells[:6]:
                    if c is not None:
                        fill = getattr(c, "fill", None)
                        if fill and getattr(fill, "fill_type", None) == "solid":
                            fg = getattr(c.fill, "fgColor", None)
                            if fg and getattr(fg, "type", None) == "rgb" and fg.rgb:
                                if _is_red_color(fg.rgb):
                                    is_red = True
                                    break
                                elif _is_yellow_color(fg.rgb):
                                    is_p3 = True
                        font = getattr(c, "font", None)
                        if font and getattr(font, "color", None):
                            fc = font.color
                            frgb = getattr(fc, "rgb", None)
                            if frgb and _is_red_color(frgb):
                                is_red = True
                                break
                # Ustunlar:
                # 0: Telefon
                # 1: Sana
                # 2: СТИР (INN)
                # 3: ЖШШИР (JSHSHIR - YaTT uchun 14 xonali)
                # 4: Корхона номи
                # 5: Онлайн-НКМ (Terminal ID / Apparat)
                # 7..18: 2025 Hisoblandi
                # 20..31: 2025 To'lov
                # 33: Qarz 2026 yil 1-yanvar
                # 34..45: 2026 Hisoblandi
                # 47..58: 2026 To'landi
                # 60: Joriy qoldiq (Колдик)
                # 62: Tuman
                
                stir_raw = row[2] if len(row) > 2 else None
                jshshir_raw = row[3] if len(row) > 3 else None
                name_raw = row[4] if len(row) > 4 else None
                terminal_raw = row[5] if len(row) > 5 else None
                phone_raw = row[0] if len(row) > 0 else None
                debt_raw = row[60] if len(row) > 60 else None
                tuman_raw = row[62] if len(row) > 62 else None
                old_debt_raw = row[33] if len(row) > 33 else 0

                stir = self.clean_digits(stir_raw)
                jshshir = self.clean_digits(jshshir_raw)
                terminal = self.clean_digits(terminal_raw)

                if not stir and not jshshir and not terminal and not name_raw:
                    continue

                # Qarz summasi
                debt_val = 0.0
                try:
                    if debt_raw is not None:
                        debt_val = float(debt_raw)
                except (ValueError, TypeError):
                    debt_val = 0.0

                old_debt = 0.0
                try:
                    if old_debt_raw is not None:
                        old_debt = float(old_debt_raw)
                except (ValueError, TypeError):
                    old_debt = 0.0

                # 2026 yil to'lovlari (col 47..58)
                payments_2026 = {}
                charges_2026 = {}
                for m_idx, m_name in enumerate(MONTHS):
                    # To'langan
                    p_col = 47 + m_idx
                    if len(row) > p_col and row[p_col]:
                        try:
                            val = float(row[p_col])
                            if val > 0:
                                payments_2026[m_name] = val
                        except (ValueError, TypeError):
                            pass
                    # Hisoblangan
                    c_col = 34 + m_idx
                    if len(row) > c_col and row[c_col]:
                        try:
                            val = float(row[c_col])
                            if val > 0:
                                charges_2026[m_name] = val
                        except (ValueError, TypeError):
                            pass

                # 2025 yil to'lovlari (col 20..31)
                payments_2025 = {}
                for m_idx, m_name in enumerate(MONTHS):
                    p_col = 20 + m_idx
                    if len(row) > p_col and row[p_col]:
                        try:
                            val = float(row[p_col])
                            if val > 0:
                                payments_2025[m_name] = val
                        except (ValueError, TypeError):
                            pass

                record = {
                    "phone": str(phone_raw or "-").strip(),
                    "stir": stir or "-",
                    "jshshir": jshshir or "-",
                    "name": str(name_raw or "Noma'lum").strip(),
                    "terminal": terminal or "-",
                    "debt": debt_val,
                    "tuman": str(tuman_raw or "-").strip(),
                    "is_p3": is_p3,
                    "is_decommissioned": is_red,
                    "device_model": "P3 Humo" if is_p3 else "Smart One",
                    "old_debt_2026": old_debt,
                    "payments_2026": payments_2026,
                    "charges_2026": charges_2026,
                    "payments_2025": payments_2025
                }

                # QIZIL BILAN BELGILANGANLAR (Hisobdan chiqarilgan / o'chirib tashlangan):
                # Faol qarz va faol apparatlar hisob-kitobiga KIRITILMAYDI!
                if is_red:
                    decom_count += 1
                    if stir:
                        new_decom_stir.setdefault(stir, []).append(record)
                    if jshshir:
                        new_decom_jshshir.setdefault(jshshir, []).append(record)
                    if terminal:
                        new_decom_terminal.setdefault(terminal, []).append(record)
                    continue

                if is_p3:
                    p3_count += 1

                if stir:
                    new_stir.setdefault(stir, []).append(record)
                if jshshir:
                    new_jshshir.setdefault(jshshir, []).append(record)
                if terminal:
                    new_terminal.setdefault(terminal, []).append(record)

                phone_clean = self.clean_digits(phone_raw)
                if phone_clean:
                    p_key = phone_clean[-9:] if len(phone_clean) >= 9 else phone_clean
                    if len(p_key) >= 7:
                        new_phone.setdefault(p_key, []).append(record)

                count += 1

            self.stir_index = new_stir
            self.jshshir_index = new_jshshir
            self.terminal_index = new_terminal
            self.phone_index = new_phone
            self.decommissioned_stir_index = new_decom_stir
            self.decommissioned_jshshir_index = new_decom_jshshir
            self.decommissioned_terminal_index = new_decom_terminal
            self.records_count = count
            self.decommissioned_count = decom_count
            wb.close()
            logger.info(
                f"Excel muvaffaqiyatli yuklandi: {count} ta faol yozuv indekslandi "
                f"(P3 Humo: {p3_count} ta, Smart One: {count - p3_count} ta, Qizil/O'chirilgan: {decom_count} ta, Phone: {len(new_phone)} ta)."
            )
            return True
        except Exception as e:
            logger.error(f"Excel faylni o'qishda xatolik: {e}")
            return False

    def update_file(self, new_file_path: str) -> bool:
        """
        Yangi kelgan Excel faylni XAVFSIZ almashtirish:
        1) Avval yangi fayl alohida nusxada sinab o'qiladi;
        2) Yozuvlar soni mantiqiy bo'lsagina latest.xlsx almashtiriladi;
        3) Aks holda eski baza o'zgarishsiz qoladi (bot bazasiz qolib ketmaydi).
        Natija sababi self.last_update_error da saqlanadi.
        """
        self.last_update_error = ""
        try:
            if new_file_path.lower().endswith(".xls"):
                self.last_update_error = "Eski .xls formati o'qilmaydi. Faylni .xlsx qilib saqlab yuboring."
                return False

            candidate = ExcelManager(excel_path=new_file_path)
            new_count = getattr(candidate, "records_count", 0)
            old_count = getattr(self, "records_count", 0)

            if new_count == 0:
                self.last_update_error = "Yangi faylda birorta ham korxona yozuvi topilmadi (format yoki ustunlar boshqacha)."
                return False
            if old_count > 0 and new_count < old_count * 0.5:
                self.last_update_error = (
                    f"Yangi faylda yozuvlar keskin kamaygan ({old_count} → {new_count}). "
                    f"Xato fayl bo'lishi mumkin, baza almashtirilmadi."
                )
                return False

            os.makedirs(DATA_DIR, exist_ok=True)
            if os.path.abspath(new_file_path) != os.path.abspath(self.excel_path):
                import shutil
                shutil.copy2(new_file_path, self.excel_path)
            return self.load_data()
        except Exception as e:
            logger.error(f"Faylni yangilashda xatolik: {e}")
            self.last_update_error = f"Faylni o'qishda xatolik: {e}"
            return False

    def search(self, query: str) -> List[Dict[str, Any]]:
        """INN (9 xonali), JSHSHIR (14 xonali), Terminal yoki Telefon raqami bo'yicha qidirish"""
        clean_q = re.sub(r"\D", "", query.strip())
        if not clean_q:
            return []

        # 1. INN (STIR) orqali tekshirish
        if clean_q in self.stir_index:
            return self.stir_index[clean_q]

        # 2. JSHSHIR orqali tekshirish
        if clean_q in self.jshshir_index:
            return self.jshshir_index[clean_q]

        # 3. Terminal raqami orqali tekshirish
        if clean_q in self.terminal_index:
            return self.terminal_index[clean_q]

        # 4. Telefon raqami orqali tekshirish (oxirgi 9 xona)
        p_key = clean_q[-9:] if len(clean_q) >= 9 else clean_q
        if p_key in self.phone_index:
            return self.phone_index[p_key]

        return []

    def check_debt(self, query: str) -> Dict[str, Any]:
        """INN / JSHSHIR bo'yicha qarz holatini tekshirish"""
        results = self.search(query)
        if not results:
            return {"found": False, "results": [], "record": None, "total_debt": 0.0, "status": "not_found"}

        first = results[0]
        total_debt = sum(item.get("debt", 0.0) for item in results)
        status = "qarzdor" if total_debt > 0 else "qarzi yo'q"
        record = {
            "org_name": first.get("name", ""),
            "inn": first.get("stir", ""),
            "status": status,
            "debt_amount": total_debt,
            "tuman": first.get("tuman", ""),
            "paid_up_to": first.get("paid_date", ""),
            "payment_date": first.get("paid_date", "")
        }
        return {
            "found": True,
            "results": results,
            "record": record,
            "total_debt": total_debt,
            "status": status
        }

    def is_decommissioned(self, query: str) -> bool:
        """
        Berilgan INN / JSHSHIR / Terminal hisobdan chiqarilgan (Excelda qizil qator) ekanligini tekshirish.
        - Agar ayni shu raqam bo'yicha kamida 1 ta faol (oq/sariq) apparat bo'lsa -> False (chunki korxona faol).
        - Agar faol apparat umuman bo'lmay, faqat qizil qator bo'lsa -> True (hisobdan chiqarilgan).
        """
        clean_q = re.sub(r"\D", "", query.strip())
        if not clean_q:
            return False
        if clean_q in self.stir_index or clean_q in self.jshshir_index or clean_q in self.terminal_index:
            return False
        return (
            clean_q in self.decommissioned_stir_index or
            clean_q in self.decommissioned_jshshir_index or
            clean_q in self.decommissioned_terminal_index
        )

    def get_decommissioned(self, query: str) -> List[Dict[str, Any]]:
        """Hisobdan chiqarilgan (qizil qator) yozuvlarni olish"""
        clean_q = re.sub(r"\D", "", query.strip())
        if not clean_q:
            return []
        if clean_q in self.decommissioned_stir_index:
            return self.decommissioned_stir_index[clean_q]
        if clean_q in self.decommissioned_jshshir_index:
            return self.decommissioned_jshshir_index[clean_q]
        if clean_q in self.decommissioned_terminal_index:
            return self.decommissioned_terminal_index[clean_q]
        return []

    def get_district_for_inn(self, inn: str) -> str:
        """Berilgan INN yoki raqam bo'yicha korxonaning tumanini aniqlash"""
        results = self.search(inn)
        if results:
            return str(results[0].get("tuman", "")).strip()
        return ""

    def is_p3_for_inn(self, query: str) -> Optional[bool]:
        """
        Berilgan INN / JSHSHIR / Terminal / Telefon bo'yicha apparat P3 Humo (sariq qator) yoki
        Smart One (oq qator) ekanligini aniqlash:
        - True: P3 Humo apparati (Excelda sariq qator)
        - False: Smart One (Excelda oq qator)
        - None: Bazadan topilmadi
        """
        if not query:
            return None
        results = self.search(query)
        if not results:
            return None
        return any(bool(r.get("is_p3", False)) for r in results)

    def get_device_model_summary(self, results: List[Dict[str, Any]]) -> str:
        """Apparatlar modellarini 'P3 Humo' va 'Smart One' bo'yicha qisqa va aniq ifodalash"""
        if not results:
            return ""
        p3_count = sum(1 for r in results if r.get("is_p3", False))
        smart_count = sum(1 for r in results if not r.get("is_p3", False))
        total = len(results)
        if total == 1:
            return "P3 Humo" if p3_count > 0 else "Smart One"
        if p3_count > 0 and smart_count > 0:
            return f"P3 Humo: {p3_count} ta, Smart One: {smart_count} ta"
        if p3_count > 0:
            return f"P3 Humo ({p3_count} ta)"
        return f"Smart One ({smart_count} ta)"

    def format_payment_history(self, results: List[Dict[str, Any]], chat_id: int = 0, user_id: int = 0) -> str:
        """
        «Mars» oylik abonent to'lovlari tarixi va hisob-kitobi.
        Faqat oylik to'lovga tegishli: qachon qaysi oy qancha to'langanini qisqa ro'yxat ko'rinishida chiqaradi (dostonsiz).
        """
        if not results:
            return (
                "⚠️ <b>To'lov tarixini tekshirish uchun korxona INN (STIR) raqamingizni yuboring.</b>\n"
                "Masalan: <code>302568138</code>"
            )

        first = results[0]
        name = first.get("name", "Noma'lum korxona")
        stir = first.get("stir", "-")
        jshshir = first.get("jshshir", "-")
        device_count = len(results)
        total_debt = sum(item.get("debt", 0.0) for item in results)

        id_str = f"<code>{stir}</code>" if stir != "-" else f"<code>{jshshir}</code>"

        MONTHS = ["Yanvar", "Fevral", "Mart", "Aprel", "May", "Iyun", "Iyul", "Avgust", "Sentyabr", "Oktyabr", "Noyabr", "Dekabr"]

        agg_charges_2026: Dict[str, float] = {}
        agg_payments_2026: Dict[str, float] = {}
        total_old_debt_2026 = 0.0

        for rec in results:
            total_old_debt_2026 += rec.get("old_debt_2026", 0.0)
            for m, s in rec.get("charges_2026", {}).items():
                agg_charges_2026[m] = agg_charges_2026.get(m, 0.0) + s
            for m, s in rec.get("payments_2026", {}).items():
                agg_payments_2026[m] = agg_payments_2026.get(m, 0.0) + s

        paid_items = []
        unpaid_items = []

        # Jami tushgan to'lov fondi (2026)
        payment_pool = sum(agg_payments_2026.values())
        effective_pool = payment_pool - max(0.0, total_old_debt_2026)

        for m in MONTHS:
            if m not in agg_charges_2026:
                if m in agg_payments_2026 and agg_payments_2026[m] > 0:
                    paid_items.append(f"• <b>{m}:</b> {agg_payments_2026[m]:,.0f} so'm ✅".replace(",", " "))
                continue

            charge = agg_charges_2026[m]
            paid_in_month = agg_payments_2026.get(m, 0.0)

            if effective_pool >= charge or paid_in_month >= charge:
                actual_paid = paid_in_month if paid_in_month > 0 else charge
                paid_items.append(f"• <b>{m}:</b> {actual_paid:,.0f} so'm ✅".replace(",", " "))
                effective_pool -= charge
            elif effective_pool > 0:
                diff = charge - effective_pool
                paid_items.append(f"• <b>{m}:</b> {effective_pool:,.0f} so'm (qisman) ⚠️".replace(",", " "))
                unpaid_items.append(f"• <b>{m}:</b> {diff:,.0f} so'm".replace(",", " "))
                effective_pool = 0
            else:
                unpaid_items.append(f"• <b>{m}:</b> {charge:,.0f} so'm".replace(",", " "))

        # Xabar matnini lo'nda va ixcham tuzish (doston yozmasdan)
        text = "📋 <b>«Mars» oylik abonent to'lovlari tarixi:</b>\n\n"
        text += f"🏢 <b>Korxona:</b> {name}\n"
        text += f"🆔 <b>СТИР (INN):</b> {id_str}"
        if device_count > 1:
            text += f" | 📟 <b>Apparatlar:</b> {device_count} ta"
        text += "\n"

        if total_debt > 0:
            text += f"🔴 <b>Joriy qarzdorlik:</b> <code>{total_debt:,.0f}</code> so'm\n\n".replace(",", " ")
        elif total_debt < 0:
            text += f"🟢 <b>Ortiqcha to'lov (avans):</b> <code>{abs(total_debt):,.0f}</code> so'm\n✅ <i>Abonent to'lovidan qarzdorlik yo'q.</i>\n\n".replace(",", " ")
        else:
            text += "🟢 <b>Abonent to'lovidan qarzdorlik yo'q (0 so'm)</b>\n\n"

        # To'langan oylar ro'yxati
        if paid_items:
            text += "🗓 <b>To'langan oylar:</b>\n"
            text += "\n".join(paid_items) + "\n\n"
        else:
            text += "🗓 <b>To'langan oylar:</b> <i>2026-yilda hali to'lov tushmagan</i>\n\n"

        # To'lanmagan oylar ro'yxati
        if unpaid_items and total_debt > 0:
            text += "❌ <b>To'lanmagan (qarzdor) oylar:</b>\n"
            text += "\n".join(unpaid_items) + "\n"

        if total_old_debt_2026 > 0:
            text += f"\n📌 <i>Eski yildan (2025) qolgan qarz: {total_old_debt_2026:,.0f} so'm</i>\n".replace(",", " ")

        if chat_id and user_id:
            from user_memory import user_memory
            last_rec = user_memory.get_last_receipt(chat_id=chat_id, user_id=user_id)
            if last_rec:
                text += "\n💡 <i>Eslatma: Oxirgi yuborgan to'lov kvitansiyangiz adminga yetkazilgan. Baza yangilangach to'lov hisobotda aks etadi.</i>\n"

        if total_debt > 0:
            text += "\n💳 <i>To'lov rekvizitlarini olish uchun «Oylik shot» deb yozing.</i>"

        return text.strip()

excel_manager = ExcelManager()
