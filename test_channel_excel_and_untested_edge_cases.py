import os
import io
import time
import shutil
import asyncio
import tempfile
import unittest
from unittest.mock import AsyncMock, MagicMock, patch
import openpyxl

from aiogram import types, Bot
from aiogram.enums import ChatType
from user_memory import UserMemoryManager, user_memory
from excel_manager import ExcelManager, excel_manager, DATA_DIR
from core.models import ReceiptFacts, FiscalChekFacts
from handlers.media import (
    handle_document,
    _process_fiscal_receipt_content,
    DESTINATION_CHANNEL_ID
)
from handlers.intents.billing import billing_intent_handler
from state_manager import get_last_file_message_id, set_last_file_message_id


class TestChannelExcelAndUntestedEdgeCases(unittest.IsolatedAsyncioTestCase):

    def setUp(self):
        user_memory.clear()

    # =========================================================================
    # 1. TEST: Kanaldan yangi to'g'ri Excel kelganda avtomatik yangilanish
    # =========================================================================
    async def test_01_channel_excel_auto_update_and_cleanup(self):
        """Kanalga yangi Excel tashlanganda: eski post o'chirilishi, baza yangilanishi"""
        # Vaqtinchalik sinov fayli yaratamiz
        wb = openpyxl.Workbook()
        ws = wb.active
        # Sarlavhalar
        ws.cell(row=1, column=1, value="TEST")
        ws.cell(row=2, column=1, value="Sarlavha")
        # 1-korxona (qator 3)
        ws.cell(row=3, column=1, value="901234567")
        ws.cell(row=3, column=3, value="309999888")
        ws.cell(row=3, column=5, value="TEST KANAL MCHJ")
        ws.cell(row=3, column=61, value=100000.0)
        ws.cell(row=3, column=63, value="Qarshi")

        tmp_dir = tempfile.mkdtemp()
        tmp_excel = os.path.join(tmp_dir, "test_channel.xlsx")
        wb.save(tmp_excel)
        wb.close()

        # Avvalgi post ID sini belgilaymiz
        set_last_file_message_id(5544)

        # Mock Bot va Message
        mock_bot = AsyncMock(spec=Bot)
        mock_bot.download = AsyncMock()
        mock_bot.delete_message = AsyncMock()

        # Download qilinganda bizning tmp_excel faylimizni temp manzilga ko'chirsin
        async def fake_download(doc, destination):
            shutil.copy2(tmp_excel, destination)
        mock_bot.download.side_effect = fake_download

        mock_doc = MagicMock(spec=types.Document)
        mock_doc.file_name = "yangilangan_baza_2026.xlsx"
        mock_doc.file_id = "doc_channel_excel_123"

        mock_chat = MagicMock(spec=types.Chat)
        mock_chat.id = DESTINATION_CHANNEL_ID
        mock_chat.type = ChatType.CHANNEL

        mock_msg = MagicMock(spec=types.Message)
        mock_msg.chat = mock_chat
        mock_msg.from_user = None  # Kanallarda from_user None bo'ladi!
        mock_msg.document = mock_doc
        mock_msg.message_id = 9988
        mock_msg.caption = "Yangi baza"
        mock_msg.media_group_id = None

        with patch.object(excel_manager, "update_file", return_value=True) as mock_update:
            await handle_document(mock_msg, mock_bot)

            # 1. update_file chaqirilishi kerak
            mock_update.assert_called_once()
            # 2. Eski post (5544) o'chirilishi kerak!
            mock_bot.delete_message.assert_called_once_with(
                chat_id=DESTINATION_CHANNEL_ID,
                message_id=5544
            )
            # 3. Yangi post ID (9988) state ga yozilgan bo'lishi kerak
            self.assertEqual(get_last_file_message_id(), 9988)

        shutil.rmtree(tmp_dir, ignore_errors=True)

    # =========================================================================
    # 2. TEST: Kanaldan buzilgan yoki yozuvlari kamaygan Excel kelganda rad etish
    # =========================================================================
    async def test_02_channel_corrupted_excel_rejected_safely(self):
        """Buzilgan Excel kelganda: eski baza saqlanib qolishi va adminga bildirishnoma borishi"""
        mock_bot = AsyncMock(spec=Bot)
        mock_bot.download = AsyncMock()

        mock_doc = MagicMock(spec=types.Document)
        mock_doc.file_name = "buzilgan_fayl.xlsx"
        mock_doc.file_id = "doc_corrupt_456"

        mock_chat = MagicMock(spec=types.Chat)
        mock_chat.id = DESTINATION_CHANNEL_ID
        mock_chat.type = ChatType.CHANNEL

        mock_msg = MagicMock(spec=types.Message)
        mock_msg.chat = mock_chat
        mock_msg.from_user = None
        mock_msg.document = mock_doc
        mock_msg.message_id = 10001
        mock_msg.caption = None
        mock_msg.media_group_id = None

        from admin_manager import admin_manager
        with patch.object(excel_manager, "update_file", return_value=False), \
             patch.object(admin_manager, "notify_admins", new_callable=AsyncMock) as mock_notify:
            excel_manager.last_update_error = "Yozuvlar keskin kamaygan (7200 -> 5)."
            await handle_document(mock_msg, mock_bot)

            # Adminga rad etilgani haqida xabarnoma yuborilishi kerak
            mock_notify.assert_called_once()
            call_text = mock_notify.call_args[1].get("text", "")
            self.assertIn("Yangi Excel qabul qilinmadi", call_text)
            self.assertIn("Yozuvlar keskin kamaygan", call_text)

    # =========================================================================
    # 3. TEST: 5 daqiqalik xotira (TTL) tugagandan keyin to'g'ri ishlash
    # =========================================================================
    async def test_03_five_minutes_memory_ttl_expiration_behavior(self):
        """5 daqiqa (300s) o'tgandan so'ng xotira tozalanishi va botning xatosiz so'rashi"""
        mem = UserMemoryManager(ttl_seconds=300)
        chat_id = -100111222333
        user_id = 777

        # 1. Foydalanuvchi hozir INN kiritdi
        mem.set_last_id(chat_id, user_id, "302568138", org_name="MARS ELEKTRON", debt=0.0)
        mem.set_last_topic(chat_id, user_id, "debt_checked")
        self.assertEqual(mem.get_last_id(chat_id, user_id), "302568138")

        # 2. Oradan 305 soniya (5 daqiqadan ortiq) vaqt o'tdi
        key = f"{chat_id}:{user_id}"
        mem.sessions[key]["last_active"] = time.time() - 305

        # 3. Xotira tozalanganda INN yo'qolishi kerak
        mem._cleanup_expired()
        self.assertIsNone(mem.get_last_id(chat_id, user_id))

        # 4. Context summary ham bo'sh bo'lishi va crash bo'lmasligi kerak
        summary = mem.get_context_summary(chat_id, user_id)
        self.assertIn("hali o'z korxona INN/JSHSHIR raqamini kiritmagan", summary)

    # =========================================================================
    # 4. TEST: Kassa chekida faqat INN bor, lekin FM yo'q holati
    # =========================================================================
    async def test_04_fiscal_chek_with_only_inn_no_fm(self):
        """Kassa chekida faqat INN o'qildi, lekin FM tushmagan bo'lsa"""
        mock_msg = MagicMock(spec=types.Message)
        mock_msg.chat = MagicMock(id=1234)
        mock_msg.from_user = MagicMock(id=5678, username="test_user")
        mock_msg.reply = AsyncMock()

        facts = ReceiptFacts(
            is_receipt=False,
            is_fiscal_chek=True,
            payer_inn="302568138",
            fiscal_module=None,
            fiscal_facts=FiscalChekFacts(
                is_fiscal_chek=True,
                inn="302568138",
                fiscal_module=None,
                org_name="MARS ELEKTRON TEXNIKA"
            )
        )

        await _process_fiscal_receipt_content(
            message=mock_msg,
            bot=AsyncMock(),
            receipt_facts=facts,
            user_script="latin",
            caption="",
            file_id="photo_only_inn_111"
        )

        mock_msg.reply.assert_called_once()
        sent_text = mock_msg.reply.call_args[0][0]
        # Mars holati chiqishi kerak
        self.assertIn("Mars", sent_text)
        # FM yo'qligi va uni yuborish so'ralishi kerak
        self.assertIn("Fiskal modul seriya raqamini", sent_text)
        # INN user_memory ga yozilgan bo'lishi kerak
        self.assertEqual(user_memory.get_last_id(1234, 5678), "302568138")
        self.assertIsNone(user_memory.get_last_fm(1234, 5678))

    # =========================================================================
    # 5. TEST: Kassa chekida faqat FM bor, lekin INN yo'q holati
    # =========================================================================
    async def test_05_fiscal_chek_with_only_fm_no_inn(self):
        """Kassa chekida faqat FM o'qildi, lekin INN tushmagan bo'lsa"""
        mock_msg = MagicMock(spec=types.Message)
        mock_msg.chat = MagicMock(id=2345)
        mock_msg.from_user = MagicMock(id=6789, username="test_user_fm")
        mock_msg.reply = AsyncMock()

        facts = ReceiptFacts(
            is_receipt=False,
            is_fiscal_chek=True,
            payer_inn=None,
            fiscal_module="LG420210630789",
            fiscal_facts=FiscalChekFacts(
                is_fiscal_chek=True,
                inn=None,
                fiscal_module="LG420210630789",
                org_name=""
            )
        )

        from ofd_bridge import ofd_bridge
        with patch.object(ofd_bridge, "query_fiscal_module", new_callable=AsyncMock) as mock_ofd:
            mock_ofd.return_value = "Litsenziya muddati: 2026-11-20 gacha faol."
            await _process_fiscal_receipt_content(
                message=mock_msg,
                bot=AsyncMock(),
                receipt_facts=facts,
                user_script="latin",
                caption="",
                file_id="photo_only_fm_222"
            )

            mock_msg.reply.assert_called_once()
            sent_text = mock_msg.reply.call_args[0][0]
            # OFD ma'lumoti chiqishi kerak
            self.assertIn("LG420210630789", sent_text)
            self.assertIn("Litsenziya muddati", sent_text)
            # Oylik to'lov uchun INN so'ralishi kerak
            self.assertIn("9 xonali INN raqamingizni yuborsangiz", sent_text)
            # FM user_memory ga yozilgan bo'lishi kerak
            self.assertEqual(user_memory.get_last_fm(2345, 6789), "LG420210630789")
            self.assertIsNone(user_memory.get_last_id(2345, 6789))

    # =========================================================================
    # 6. TEST: Telefon raqamlari turli formatda yozilganda qidiruv
    # =========================================================================
    def test_06_phone_search_various_formats(self):
        """Telefon raqami +998938900990, 93 890 09 90 yoki 93-890-09-90 yozilganda topilishi"""
        formats = [
            "938900990",
            "+998938900990",
            "998938900990",
            "+998 93 890 09 90",
            "93-890-09-90"
        ]
        for p in formats:
            clean_digits = excel_manager.clean_digits(p)
            q_key = clean_digits[-9:] if len(clean_digits) >= 9 else clean_digits
            res = excel_manager.phone_index.get(q_key, [])
            self.assertTrue(len(res) > 0, f"Telefon formati topilmadi: {p} (q_key={q_key})")
            self.assertEqual(res[0].get("stir"), "311790999")

    # =========================================================================
    # 7. TEST: Guruhda parallel 3 ta har xil media (Chek, Kvitansiya, Xato rasm)
    # =========================================================================
    async def test_07_group_chat_mixed_media_isolation(self):
        """Umumiy guruhda User A (fiskal chek), User B (to'lov), User C (apparat) bir vaqtda kirishi"""
        group_id = -100999000111
        user_a = 101  # Kassa cheki
        user_b = 102  # Bank to'lovi
        user_c = 103  # Apparat ekrani

        # User A kassa cheki (INN 308123456)
        user_memory.set_last_id(group_id, user_a, "308123456")
        user_memory.set_last_fm(group_id, user_a, "LG420111")
        user_memory.set_last_topic(group_id, user_a, "fiscal_receipt_checked")

        # User B bank to'lovi (INN 302568138, 100000 so'm)
        user_memory.set_last_id(group_id, user_b, "302568138")
        user_memory.set_last_receipt(group_id, user_b, {"amount": 100000.0, "debt_covered": True})

        # User C apparat ekrani (nosozlik)
        user_memory.set_last_topic(group_id, user_c, "pax_d210_menu")

        # Tekshiruvlar: hech kimning ma'lumotlari aralashib ketmasligi shart!
        self.assertEqual(user_memory.get_last_id(group_id, user_a), "308123456")
        self.assertEqual(user_memory.get_last_fm(group_id, user_a), "LG420111")
        self.assertIsNone(user_memory.get_last_receipt(group_id, user_a))

        self.assertEqual(user_memory.get_last_id(group_id, user_b), "302568138")
        self.assertIsNotNone(user_memory.get_last_receipt(group_id, user_b))
        self.assertIsNone(user_memory.get_last_fm(group_id, user_b))

        self.assertIsNone(user_memory.get_last_id(group_id, user_c))
        self.assertEqual(user_memory.get_last_topic(group_id, user_c), "pax_d210_menu")

    # =========================================================================
    # 8. TEST: Shot raqam va rekvizitlar so'rovlari
    # =========================================================================
    async def test_08_payment_requisites_inquiries(self):
        """Mijoz 'shot raqam tashlang', 'yillik schet', 'oylik shot' deb so'raganda to'g'ri ajratilishi"""
        from payment_accounts import classify_payment_intent

        # 1. Yillik to'lov hisob raqami
        intent_annual = classify_payment_intent("yillik tolov uchun shot raqam tashvorin")
        self.assertEqual(intent_annual, "annual")

        # 2. Oylik to'lov hisob raqami
        intent_monthly = classify_payment_intent("oylik tolov scheti qaysi edi")
        self.assertEqual(intent_monthly, "monthly")

        # 3. Umumiy to'lov yoki ikkalasi
        intent_both = classify_payment_intent("to'lov qilish uchun rekvizitlar bering")
        self.assertEqual(intent_both, "both")

    # =========================================================================
    # 9. TEST: Kirill yozuvidagi qarz va to'lov savollari
    # =========================================================================
    async def test_09_cyrillic_script_handling_and_debt_query(self):
        """Kirill yozuvida 'қарзим борми' yoki 'йиллик тўлов' deb yozganda javob kirillcha chiqishi"""
        from translit_manager import detect_script, adapt_to_user_script

        text_cyr = "Салом, бизнинг қарзимиз борми 302568138"
        script = detect_script(text_cyr)
        self.assertEqual(script, "cyrillic")

        sample_reply = "✅ «Mars» oylik abonent to'lovidan qarzdorlik yo'q (0 so'm)."
        adapted = adapt_to_user_script(sample_reply, script)
        self.assertIn("«Марс»", adapted)
        self.assertIn("қарздорлик йўқ", adapted)

    # =========================================================================
    # 10. TEST: Xabarda bir nechta INN kelganda
    # =========================================================================
    def test_10_multiple_inns_in_single_message(self):
        """Bitta xabarda ikkita INN kelganda xavfsiz va aniq ajratilishi"""
        import re
        msg = "Bizda ikkita korxona bor: 302568138 va 308123456 qaysi birida qarz bor?"
        inns = re.findall(r"\b(\d{9})\b", msg)
        self.assertEqual(len(inns), 2)
        self.assertIn("302568138", inns)
        self.assertIn("308123456", inns)


if __name__ == "__main__":
    unittest.main()
